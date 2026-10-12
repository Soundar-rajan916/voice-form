import pytest
from httpx import AsyncClient

@pytest.mark.asyncio
async def test_admin_create_and_publish_form(client: AsyncClient, admin_token_headers):
    # 1. Create form
    payload = {
        "title": "Citizen Feedback Form",
        "description": "Voice survey for citizens",
        "fields": [
            {
                "field_name": "citizen_name",
                "data_type": "String",
                "description": "Full name of the citizen",
                "required": True,
                "position": 0
            },
            {
                "field_name": "rating",
                "data_type": "Number",
                "description": "Rating from 1 to 5",
                "required": True,
                "position": 1
            }
        ]
    }
    res = await client.post("/admin/forms", json=payload, headers=admin_token_headers)
    assert res.status_code == 201
    form_data = res.json()
    form_id = form_data["id"]
    assert form_data["status"] == "draft"
    assert form_data["current_version"]["version_number"] == 1
    assert len(form_data["current_version"]["fields"]) == 2

    # 2. Publish form
    pub_res = await client.post(f"/admin/forms/{form_id}/publish", headers=admin_token_headers)
    assert pub_res.status_code == 200
    pub_data = pub_res.json()
    assert pub_data["status"] == "published"
    fields = pub_data["current_version"]["fields"]
    for f in fields:
        assert f["question_en"] is not None and len(f["question_en"]) > 0
        assert f["question_ta"] is not None and len(f["question_ta"]) > 0

@pytest.mark.asyncio
async def test_user_cannot_access_admin_endpoints(client: AsyncClient, user_token_headers):
    payload = {
        "title": "Unauthorized Form",
        "description": "Test",
        "fields": []
    }
    res = await client.post("/admin/forms", json=payload, headers=user_token_headers)
    assert res.status_code == 403 # Forbidden

@pytest.mark.asyncio
async def test_form_versioning_and_submission_isolation(
    client: AsyncClient,
    admin_token_headers,
    user_token_headers
):
    # 1. Admin creates and publishes Form V1
    create_res = await client.post(
        "/admin/forms",
        json={
            "title": "Customer Feedback",
            "fields": [
                {"field_name": "name", "data_type": "String", "description": "name"}
            ]
        },
        headers=admin_token_headers
    )
    form_id = create_res.json()["id"]
    await client.post(f"/admin/forms/{form_id}/publish", headers=admin_token_headers)

    # 2. User starts submission S1 on V1
    sub1_res = await client.post(f"/forms/{form_id}/submissions", headers=user_token_headers)
    assert sub1_res.status_code == 201
    s1_data = sub1_res.json()
    s1_id = s1_data["id"]
    assert s1_data["version_number"] == 1

    # 3. Admin edits form (since it has submissions, this must create Version 2)
    update_res = await client.put(
        f"/admin/forms/{form_id}",
        json={
            "title": "Customer Feedback Updated",
            "fields": [
                {"field_name": "name", "data_type": "String", "description": "name"},
                {"field_name": "age", "data_type": "Number", "description": "age in years"}
            ]
        },
        headers=admin_token_headers
    )
    assert update_res.status_code == 200
    updated_data = update_res.json()
    assert updated_data["current_version"]["version_number"] == 2
    assert len(updated_data["current_version"]["fields"]) == 2

    # Publish V2
    await client.post(f"/admin/forms/{form_id}/publish", headers=admin_token_headers)

    # 4. User starts submission S2 (gets V2)
    sub2_res = await client.post(f"/forms/{form_id}/submissions", headers=user_token_headers)
    s2_data = sub2_res.json()
    assert s2_data["version_number"] == 2
    assert len(s2_data["fields"]) == 2

    # 5. S1 should still show V1 (with 1 field)
    s1_get = await client.get(f"/submissions/{s1_id}", headers=user_token_headers)
    assert s1_get.status_code == 200
    assert s1_get.json()["version_number"] == 1
    assert len(s1_get.json()["fields"]) == 1

@pytest.mark.asyncio
async def test_other_user_submission_returns_404(
    client: AsyncClient,
    admin_token_headers,
    user_token_headers
):
    # Register user2
    await client.post("/auth/register", json={"email": "user2@test.com", "password": "User2Pass"})
    login2 = await client.post("/auth/login", json={"email": "user2@test.com", "password": "User2Pass"})
    user2_token = login2.json()["access_token"]
    user2_headers = {"Authorization": f"Bearer {user2_token}"}

    # Admin creates & publishes form
    form_res = await client.post("/admin/forms", json={"title": "Survey", "fields": [{"field_name": "x", "data_type": "String", "description": "x"}]}, headers=admin_token_headers)
    form_id = form_res.json()["id"]
    await client.post(f"/admin/forms/{form_id}/publish", headers=admin_token_headers)

    # User 1 creates submission
    s1_res = await client.post(f"/forms/{form_id}/submissions", headers=user_token_headers)
    s1_id = s1_res.json()["id"]

    # User 2 tries to read S1 -> must get 404 (NOT 403)
    get_res = await client.get(f"/submissions/{s1_id}", headers=user2_headers)
    assert get_res.status_code == 404
    assert get_res.json()["detail"] == "Submission not found"

    # User 2 tries to delete S1 -> must get 404 (NOT 403)
    del_res = await client.delete(f"/submissions/{s1_id}", headers=user2_headers)
    assert del_res.status_code == 404

@pytest.mark.asyncio
async def test_admin_archive_and_unarchive_form(client: AsyncClient, admin_token_headers):
    # 1. Create and publish form
    payload = {
        "title": "Archive Test Form",
        "description": "Test archiving and unarchiving",
        "fields": [{"field_name": "name", "data_type": "String", "description": "Applicant Name"}]
    }
    form_res = await client.post("/admin/forms", json=payload, headers=admin_token_headers)
    assert form_res.status_code == 201
    form_id = form_res.json()["id"]

    pub_res = await client.post(f"/admin/forms/{form_id}/publish", headers=admin_token_headers)
    assert pub_res.status_code == 200
    assert pub_res.json()["status"] == "published"

    # 2. Archive form
    arch_res = await client.post(f"/admin/forms/{form_id}/archive", headers=admin_token_headers)
    assert arch_res.status_code == 200
    assert arch_res.json()["status"] == "archived"

    # 3. Unarchive form (restores to published since it has fields)
    unarch_res = await client.post(f"/admin/forms/{form_id}/unarchive", headers=admin_token_headers)
    assert unarch_res.status_code == 200
    assert unarch_res.json()["status"] == "published"

    # 4. Archive again and unarchive explicitly as draft
    await client.post(f"/admin/forms/{form_id}/archive", headers=admin_token_headers)
    unarch_draft = await client.post(f"/admin/forms/{form_id}/unarchive?target_status=draft", headers=admin_token_headers)
    assert unarch_draft.status_code == 200
    assert unarch_draft.json()["status"] == "draft"

@pytest.mark.asyncio
async def test_admin_export_excel(client: AsyncClient, admin_token_headers, user_token_headers):
    import io
    import openpyxl

    # 1. Create and publish form
    form_res = await client.post(
        "/admin/forms",
        json={
            "title": "Feedback Excel Form",
            "description": "Form for testing excel export",
            "fields": [
                {"field_name": "full_name", "data_type": "String", "description": "Full Name", "required": True, "position": 0},
                {"field_name": "age", "data_type": "Number", "description": "Age", "required": False, "position": 1}
            ]
        },
        headers=admin_token_headers
    )
    assert form_res.status_code == 201
    form_id = form_res.json()["id"]
    await client.post(f"/admin/forms/{form_id}/publish", headers=admin_token_headers)

    # 2. User creates an in_progress submission (should NOT be exported)
    in_prog = await client.post(f"/forms/{form_id}/submissions", headers=user_token_headers)
    assert in_prog.status_code == 201

    # 3. User creates a completed submission
    comp = await client.post(f"/forms/{form_id}/submissions", headers=user_token_headers)
    comp_id = comp.json()["id"]
    sub_detail = comp.json()
    fields = sub_detail["fields"]
    name_field = next(f for f in fields if f["field_name"] == "full_name")

    # Submit with answer
    submit_res = await client.post(
        f"/submissions/{comp_id}/submit",
        json=[
            {"field_id": name_field["id"], "value": "John Doe", "raw_transcript": "My name is John Doe", "language": "en"}
        ],
        headers=user_token_headers
    )
    assert submit_res.status_code == 200

    # 4. Request Excel export via Bearer token
    res = await client.get(f"/admin/forms/{form_id}/submissions.xlsx", headers=admin_token_headers)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=" in res.headers["content-disposition"]

    # 5. Inspect Excel workbook contents using openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Submissions" in wb.sheetnames
    assert "Answers Ledger" in wb.sheetnames

    sub_sheet = wb["Submissions"]
    # Header row is row 1
    headers = [cell.value for cell in sub_sheet[1]]
    assert "Reference Number" not in headers
    assert "Channel" not in headers
    assert "Applicant Email" in headers
    assert "Full Name" in headers

    # Should have exactly 2 rows (header + 1 submitted row, in-progress excluded)
    rows = list(sub_sheet.iter_rows(values_only=True))
    assert len(rows) == 2
    row_data = dict(zip(rows[0], rows[1]))
    assert row_data["Full Name"] == "John Doe"

    # Also test export with token in query string (for direct browser link)
    admin_token = admin_token_headers["Authorization"].split(" ")[1]
    res_query = await client.get(f"/admin/forms/{form_id}/submissions.xlsx?token={admin_token}")
    assert res_query.status_code == 200

